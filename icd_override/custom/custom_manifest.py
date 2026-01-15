import frappe
from openpyxl import load_workbook
from datetime import datetime
from icd_tz.icd_tz.doctype.manifest.manifest import Manifest

class CustomManifest(Manifest):
    @frappe.whitelist()
    def extract_data_from_manifest_file(self):
        #frappe.msgprint("ICD_TZ  OVERRIDE- Extracting data from manifest file...")
        if self.manifest:
            file_url = self.manifest
            file_path = frappe.get_site_path('private', 'files', file_url.split('/files/')[-1])
            
            # Load the Excel file
            workbook = load_workbook(file_path, data_only=True)

            # Function to convert dates
            def convert_date(excel_date):
                if isinstance(excel_date, datetime):
                    return excel_date.strftime('%Y-%m-%d')
                if isinstance(excel_date, str):
                    try:
                        return datetime.strptime(excel_date, '%d/%m/%Y').strftime('%Y-%m-%d')
                    except ValueError:
                        return None
                return None

            # Process the MRN Detail (1) sheet
            vessel_info_sheet = workbook['MRN Detail (1)']
            vessel_info_row = next(vessel_info_sheet.iter_rows(min_row=4, values_only=True))
            self.mrn = vessel_info_row[0]
            self.vessel_name = vessel_info_row[1]
            self.call_sign = vessel_info_row[2]
            self.voyage_no = vessel_info_row[3]
            self.custom_departure_date = convert_date(vessel_info_row[4])        
            self.arrival_date = convert_date(vessel_info_row[5])
            self.tpa_uid = vessel_info_row[6]
            
            # Process the Master BL List (4) sheet
            master_bl_sheet = workbook['Master BL List (4)']            
            self.master_bl = []
            filtered_rows_data = []
            filtered_HBL_rows_data = []
            PMMICD = "WITZDL034" #PMM ICD number
            consolidate_container = "C"  # House BL type
            for row in master_bl_sheet.iter_rows(min_row=4, values_only=True):                
                if row[4].strip().upper() == PMMICD:
                    filtered_rows_data.append(row)
                if row[4].strip().upper() == PMMICD and row[2] == consolidate_container.strip().upper():
                    filtered_HBL_rows_data.append(row)

            # Process the Container (2) sheet
            containers_sheet = workbook['Container (2)']
            self.populate_containers(filtered_rows_data, containers_sheet)                                      
            
            # Process the HBL Container (3) sheet
            hbl_containers_sheet = workbook['HBL Container (3)']
            self.populate_hbl_containers(filtered_rows_data, hbl_containers_sheet) 

            # Process the Master BL List (4) sheet
            self.populate_masterbl(filtered_rows_data)

            # Process the House BL List (5) sheet
            house_bl_sheet = workbook['House BL List (5)']
            self.populate_house_bl_containers(filtered_HBL_rows_data, house_bl_sheet)
            # self.save()
            return False

    def populate_house_bl_containers(self, filtered_HBL_rows_data, house_bl_sheet):
        self.house_bl = []
        for row in filtered_HBL_rows_data:
            target_value = row[0].strip().lower()
            for house_bl_row in house_bl_sheet.iter_rows(min_row=4, values_only=True):
                if (target_value == house_bl_row[0].strip().lower()):
                    house_bl = self.append("house_bl", {})
                    house_bl.m_bl_no = house_bl_row[0]
                    house_bl.h_bl_no = house_bl_row[1]
                    house_bl.cargo_classification = house_bl_row[2]
                    house_bl.place_of_destination = house_bl_row[3]
                    house_bl.net_weight = house_bl_row[4]
                    house_bl.net_weight_unit = house_bl_row[5]
                    house_bl.number_of_containers = house_bl_row[6]
                    house_bl.description_of_goods = house_bl_row[7]
                    house_bl.number_of_package = house_bl_row[8]
                    house_bl.package_unit = house_bl_row[9]
                    house_bl.gross_weight = house_bl_row[10]
                    house_bl.gross_weight_unit = house_bl_row[11]
                    house_bl.gross_volume = house_bl_row[12]
                    house_bl.gross_volume_unit = house_bl_row[13]
                    house_bl.invoice_value = house_bl_row[14]
                    house_bl.invoice_currency = house_bl_row[15]
                    house_bl.freight_charge = house_bl_row[16]
                    house_bl.freight_currency = house_bl_row[17]
                    house_bl.imdg_code = house_bl_row[18]
                    house_bl.packing_type = house_bl_row[19]
                    house_bl.shipping_agent_code = house_bl_row[20]
                    house_bl.shipping_agent_name = house_bl_row[21]
                    house_bl.forwarder_code = house_bl_row[22]
                    house_bl.forwarder_name = house_bl_row[23]
                    house_bl.exporter_name = house_bl_row[24]
                    house_bl.exporter_tel = house_bl_row[25]
                    house_bl.exporter_address = house_bl_row[26]
                    house_bl.exporter_tin = house_bl_row[27]
                    house_bl.consignee_name = house_bl_row[28]
                    house_bl.consignee_tel = house_bl_row[29]
                    house_bl.consignee_address = house_bl_row[30]
                    house_bl.consignee_tin = house_bl_row[31]
                    house_bl.notify_name = house_bl_row[32]
                    house_bl.notify_tel = house_bl_row[33]
                    house_bl.notify_address = house_bl_row[34]
                    house_bl.notify_tin = house_bl_row[35]
                    house_bl.shipping_mark = house_bl_row[36]
                    house_bl.oil_type = house_bl_row[37]

    def populate_hbl_containers(self, filtered_rows_data, hbl_containers_sheet):
        self.hbl_containers = []
        for row in filtered_rows_data:
            target_value = row[0].strip().lower() 
            for hbl_container_row in hbl_containers_sheet.iter_rows(min_row=4, values_only=True):
                if (target_value == hbl_container_row[0].strip().lower()):
                    hbicontainer = self.append("hbl_containers", {})
                    hbicontainer.m_bl_no = hbl_container_row[0]
                    hbicontainer.h_bl_no = hbl_container_row[1]
                    hbicontainer.type_of_container = hbl_container_row[2]
                    hbicontainer.container_no = hbl_container_row[3]
                    hbicontainer.container_size = hbl_container_row[4]
                    hbicontainer.seal_no1 = hbl_container_row[5]
                    hbicontainer.seal_no2 = hbl_container_row[6]
                    hbicontainer.seal_no3 = hbl_container_row[7]
                    hbicontainer.freight_indicator = hbl_container_row[8]
                    hbicontainer.no_of_packages = hbl_container_row[9]
                    hbicontainer.package_unit = hbl_container_row[10]    
                    hbicontainer.weight = hbl_container_row[11]
                    hbicontainer.weight_unit = hbl_container_row[12]
                    hbicontainer.plug_type_of_reefer = hbl_container_row[13]
                    hbicontainer.minimum_temperature = hbl_container_row[14]
                    hbicontainer.maximum_temperature = hbl_container_row[15]

    def populate_masterbl(self, filtered_rows_data):
        for row in filtered_rows_data:
                master_bl = self.append("master_bl", {})
                master_bl.m_bl_no = row[0]
                master_bl.cargo_classification = row[1]
                master_bl.bl_type = row[2]
                master_bl.place_of_destination = row[3]
                master_bl.place_of_delivery = row[4]
                master_bl.oil_type = row[5]
                master_bl.port_of_loading = row[6]
                master_bl.number_of_containers = row[7]
                master_bl.cargo_description = row[8]
                master_bl.number_of_package = row[9]
                master_bl.package_unit = row[10]
                master_bl.gross_weight = row[11]
                master_bl.gross_weight_unit = row[12]
                master_bl.gross_volume = row[13]
                master_bl.gross_volume_unit = row[14]
                master_bl.invoice_value = row[15]
                master_bl.invoice_currency = row[16]
                master_bl.freight_charge = row[17]
                master_bl.freight_currency = row[18]
                master_bl.imdg_code = row[19]
                master_bl.packing_type = row[20]
                master_bl.shipping_agent_code = row[21]
                master_bl.shipping_agent_name = row[22]
                master_bl.forwarder_code = row[23]
                master_bl.forwarder_name = row[24]
                master_bl.forwarder_tel = row[25]
                master_bl.exporter_name = row[26]
                master_bl.exporter_tel = row[27]
                master_bl.exporter_address = row[28]
                master_bl.exporter_tin = row[29]
                master_bl.consignee_name = row[30]
                master_bl.consignee_tel = row[31]
                master_bl.consignee_address = row[32]
                master_bl.consignee_tin = row[33]
                master_bl.notify_name = row[34]
                master_bl.notify_tel = row[35]
                master_bl.notify_address = row[36]
                master_bl.notify_tin = row[37]
                master_bl.shipping_mark = row[38]
                master_bl.net_weight = row[39]
                master_bl.net_weight_unit = row[40]

    def populate_containers(self, filtered_rows_data, containers_sheet):
        self.containers = []
        for row in filtered_rows_data:
            target_value = row[0].strip().lower()               
            for container_row in containers_sheet.iter_rows(min_row=4, values_only=True):
                if (target_value == container_row[0].strip().lower()):
                    container = self.append("containers", {})
                    container.m_bl_no = container_row[0]
                    container.type_of_container = container_row[1]
                    container.container_no = container_row[2]
                    container.container_size = container_row[3]
                    container.seal_no1 = container_row[4]
                    container.seal_no2 = container_row[5]
                    container.seal_no3 = container_row[6]
                    container.freight_indicator = container_row[7]
                    container.no_of_packages = container_row[8]
                    container.package_unit = container_row[9]                    
                    container.weight = container_row[10]
                    container.weight_unit = container_row[11]
                    container.plug_type_of_reefer = container_row[12]
                    container.minimum_temperature = container_row[13]
                    container.maximum_temperature = container_row[14]